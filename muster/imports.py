import csv
import io
import ipaddress
import json
import math
from datetime import date, datetime, time
from urllib.parse import urlsplit, urlunsplit

from openpyxl import load_workbook

MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024


def normalize_url(value):
    value = str(value or "").strip()
    try:
        parts = urlsplit(value)
        host = parts.hostname
        if (
            parts.scheme not in ("http", "https")
            or not host
            or parts.username
            or parts.password
            or any(c.isspace() or ord(c) < 32 for c in value)
        ):
            return None
        # Collection is for public web pages; reject obvious local targets up front.
        if host.lower() == "localhost" or host.lower().endswith((".localhost", ".local")):
            return None
        try:
            if not ipaddress.ip_address(host).is_global:
                return None
        except ValueError:
            pass
        port = parts.port
        netloc = host.lower() if ":" not in host else f"[{host.lower()}]"
        if port and (parts.scheme, port) not in (("http", 80), ("https", 443)):
            netloc += f":{port}"
        return urlunsplit((parts.scheme, netloc, parts.path or "/", parts.query, ""))
    except ValueError:
        return None


def parse_import(filename: str, content: bytes, pasted: str):
    if len(content) > MAX_BYTES:
        raise ValueError("Choose a spreadsheet smaller than 5 MB.")
    cell_kinds = {}
    headers = []
    if filename:
        if filename.lower().endswith(".csv"):
            try:
                rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig")), strict=True))
            except (UnicodeError, csv.Error) as exc:
                raise ValueError(
                    "Could not read the CSV. Save it as UTF-8 with a url column."
                ) from exc
        elif filename.lower().endswith(".xlsx"):
            workbook = None
            try:
                # Limit expanded archive size before openpyxl reads shared strings.
                from zipfile import ZipFile

                with ZipFile(io.BytesIO(content)) as archive:
                    if sum(entry.file_size for entry in archive.infolist()) > 25 * 1024 * 1024:
                        raise ValueError("The expanded spreadsheet is too large.")
                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
                rows = []
                for row_index, row in enumerate(workbook.active.iter_rows()):
                    rows.append([cell.value for cell in row])
                    for column_index, cell in enumerate(row):
                        if cell.data_type == "f":
                            cell_kinds[(row_index, column_index)] = "formula"
                    if len(rows) > MAX_ROWS + 1:
                        raise ValueError("Import at most 5,000 rows at a time.")
            except Exception as exc:
                raise ValueError(
                    "Could not read the XLSX. Use a valid spreadsheet with a url column (at most 5,000 rows)."
                ) from exc
            finally:
                if workbook:
                    workbook.close()
        else:
            raise ValueError("Choose a .csv or .xlsx file.")
        if not rows:
            raise ValueError("Your spreadsheet needs a column named url.")
        headers = [str(cell or "").strip() for cell in rows[0]]
        names = [header.casefold() for header in headers]
        if names.count("url") != 1:
            raise ValueError("Your spreadsheet needs exactly one column named url.")
        if (
            len(headers) > 51
            or len(set(names)) != len(names)
            or any(len(header) > 100 for header in headers)
        ):
            raise ValueError(
                "Use unique column names of at most 100 characters, with up to 50 columns besides url."
            )
        column = names.index("url")
        if any(len(row) > len(headers) for row in rows[1:]):
            raise ValueError("Some rows have more cells than the header. Give every column a name.")
        if any(
            not headers[i]
            and any(i < len(row) and row[i] is not None and row[i] != "" for row in rows[1:])
            for i in range(len(headers))
        ):
            raise ValueError("Give every column containing values a name.")
        values = [row[column] if len(row) > column else "" for row in rows[1:]]
    else:
        values = pasted.splitlines()
    if len(values) > MAX_ROWS:
        raise ValueError("Import at most 5,000 rows at a time.")
    urls, invalid, duplicates = [], [], 0
    seen = {}
    records = []
    feature_bytes = 0
    for index, value in enumerate(values, start=1):
        url = normalize_url(value)
        if not url:
            invalid.append(index)
        else:
            input_row = None
            if headers:
                cells, kinds = {}, {}
                row = rows[index]
                for position, header in enumerate(headers):
                    if position == column or not header:
                        continue
                    value = row[position] if position < len(row) else None
                    kind = cell_kinds.get((index, position))
                    if isinstance(value, (datetime, date, time)):
                        value, kind = value.isoformat(), "date"
                    elif isinstance(value, bool):
                        kind = "boolean"
                    elif isinstance(value, (int, float)):
                        if isinstance(value, float) and not math.isfinite(value):
                            raise ValueError("Spreadsheet numbers must be finite.")
                        kind = "number"
                    elif value is None:
                        kind = "empty"
                    else:
                        value = str(value)
                        kind = kind or "text"
                    if isinstance(value, str) and len(value) > 10000:
                        raise ValueError("Keep each spreadsheet cell within 10,000 characters.")
                    cells[header], kinds[header] = value, kind
                input_row = {
                    "row": index,
                    "original_url": str(values[index - 1]),
                    "cells": cells,
                    "kinds": kinds,
                }
                feature_bytes += len(json.dumps(input_row, ensure_ascii=False).encode("utf-8"))
                if feature_bytes > MAX_BYTES:
                    raise ValueError(
                        "The retained spreadsheet values exceed 5 MB. Import a smaller spreadsheet."
                    )
            if url in seen:
                duplicates += 1
                if input_row:
                    records[seen[url]]["input_rows"].append(input_row)
            else:
                seen[url] = len(records)
                urls.append(url)
                records.append({"url": url, "input_rows": [input_row] if input_row else []})
    if not values:
        raise ValueError("Add at least one link or choose a spreadsheet.")
    return {
        "urls": urls,
        "invalid": invalid,
        "duplicates": duplicates,
        "total": len(values),
        "records": records,
    }
