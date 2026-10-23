"""Bounded parsers. Source bytes are never rewritten, evaluated or executed."""

import bisect
import csv
import hashlib
import io
import math
import re
import unicodedata
import zipfile
from array import array
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import PurePosixPath
from uuid import NAMESPACE_URL, UUID, uuid5

from defusedxml import ElementTree
from openpyxl import load_workbook

from quartermaster.contracts.sources import Chunk, FieldSummary, ParseOptions, SheetPreview
from quartermaster.errors import DomainError

MAX_UPLOAD = 8 * 1024 * 1024
MAX_EXPANDED = 32 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 100
MAX_CELLS = 100_000
MAX_CELL_LENGTH = 8000
CHUNK_LENGTH = 1500
MAX_HEADINGS = 2000


@dataclass
class Row:
    number: int
    line_start: int
    line_end: int
    values: dict[str, str | None]
    formulas: set[str]


@dataclass
class Table:
    name: str
    fields: list[str]
    rows: list[Row]


@dataclass
class Parsed:
    tables: list[Table]
    markdown: str | None = None


def invalid(message: str):
    raise DomainError(f"SOURCE_INVALID: {message}")


def safe_name(filename: str) -> tuple[str, str]:
    leaf = filename.replace("\\", "/").rsplit("/", 1)[-1]
    leaf = "".join(c for c in leaf if not unicodedata.category(c).startswith("C"))
    kind = PurePosixPath(leaf).suffix.lower().lstrip(".")
    if kind not in {"csv", "xlsx", "md"}:
        invalid("Choose a .csv, .xlsx or .md file.")
    # Display name only; it is never part of a filesystem path or download header.
    return leaf[:180] or f"upload.{kind}", kind


def decode(data: bytes, encoding: str) -> str:
    try:
        text = data.decode(encoding, errors="strict")
    except UnicodeError:
        invalid("Choose the correct encoding (UTF-8 or Windows-1252) and upload again.")
    if any(ord(c) < 32 and c not in "\t\r\n" for c in text) or "\x7f" in text:
        invalid("Text files must contain text, without binary control characters.")
    return text


def headers(values: list[str | None]) -> list[str]:
    if not values or len(values) > MAX_COLUMNS:
        invalid("Use between 1 and 100 columns, with a field name in the first row.")
    if any(v is None or not v.strip() or len(v) > 200 for v in values):
        invalid("Every first-row field name must contain 1–200 characters.")
    if len(set(values)) != len(values):
        invalid("Give each field a unique first-row name before uploading.")
    return values


def check_table(table: Table):
    if len(table.rows) > MAX_ROWS:
        invalid("Split the file into tables of at most 10,000 data rows.")
    if any(
        len(value or "") > MAX_CELL_LENGTH for row in table.rows for value in row.values.values()
    ):
        invalid("Keep each cell below 8,001 characters.")


def parse_csv(data: bytes, options: ParseOptions) -> Parsed:
    text = decode(data, options.encoding)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=options.delimiter, strict=True)
    try:
        fields = headers(next(reader, []))
        rows = []
        previous_line = reader.line_num
        for number, values in enumerate(reader, start=2):
            start = previous_line + 1
            previous_line = reader.line_num
            if len(values) != len(fields):
                invalid("Each CSV record must have the same number of fields as its header.")
            rows.append(
                Row(number, start, reader.line_num, dict(zip(fields, values, strict=True)), set())
            )
            if len(rows) > MAX_ROWS or len(rows) * len(fields) > MAX_CELLS:
                invalid("Split the CSV into at most 10,000 rows and 100,000 cells.")
    except csv.Error:
        invalid("Check CSV quoting and the selected delimiter; the CSV could not be read.")
    table = Table("CSV", fields, rows)
    check_table(table)
    return Parsed([table])


def cell_text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and not math.isfinite(value):
        invalid("Remove non-finite numeric cells before uploading.")
    return str(value)


def inspect_zip(data: bytes):
    # Validate expanded bytes and every XML document before openpyxl sees the archive.
    # No archive entry is extracted to disk, and no external relationship is followed.
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        infos = archive.infolist()
        if len(infos) > 512 or len({i.filename for i in infos}) != len(infos):
            invalid("Use a workbook with at most 512 unique archive entries.")
        if sum(i.file_size for i in infos) > MAX_EXPANDED:
            invalid("Use a workbook smaller than 32 MiB when expanded.")
        for info in infos:
            path = PurePosixPath(info.filename)
            lower = info.filename.lower()
            if path.is_absolute() or ".." in path.parts or info.flag_bits & 1:
                invalid("Encrypted or unsafe workbook archives are unsupported.")
            if any(part in lower for part in ("vbaproject", "macrosheets/", "embeddings/")):
                invalid("Remove macros and embedded objects; only ordinary XLSX is accepted.")
            if lower.endswith((".xml", ".rels")):
                xml = archive.read(info)
                root = ElementTree.fromstring(xml, forbid_dtd=True)
                if lower == "[content_types].xml" and any(
                    "macroenabled" in str(node.attrib).lower() for node in root
                ):
                    invalid("Macro-enabled workbooks are unsupported.")
                if lower.startswith("xl/worksheets/") and lower.endswith(".xml"):
                    for node in root.iter():
                        if node.tag.rsplit("}", 1)[-1] == "c":
                            ref = node.attrib.get("r", "")
                            match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", ref)
                            if not match:
                                invalid("The workbook has an invalid cell address.")
                            col = 0
                            for letter in match[1]:
                                col = col * 26 + ord(letter) - 64
                            if col > MAX_COLUMNS or int(match[2]) > MAX_ROWS + 1:
                                invalid("Keep worksheets within 100 columns and 10,001 total rows.")


def parse_xlsx(data: bytes) -> Parsed:
    try:
        inspect_zip(data)
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        try:
            if len(workbook.worksheets) > 20:
                invalid("Use a workbook with at most 20 worksheets.")
            tables = []
            count = 0
            for sheet in workbook.worksheets:
                # Ignore unreliable stored dimensions. Addresses were bounded above.
                sheet.reset_dimensions()
                iterator = iter(sheet.iter_rows())
                first = next(iterator, ())
                if not first:
                    tables.append(Table(sheet.title, [], []))
                    continue
                if any(cell.data_type == "f" for cell in first):
                    invalid("Use literal field names rather than formulas in the header.")
                fields = headers([cell_text(c.value) for c in first])
                rows = []
                for number, cells in enumerate(iterator, start=2):
                    if len(cells) > len(fields):
                        invalid("Every populated worksheet column needs a first-row field name.")
                    values = [cell_text(cell.value) for cell in cells]
                    values.extend([None] * (len(fields) - len(values)))
                    formulas = {fields[i] for i, cell in enumerate(cells) if cell.data_type == "f"}
                    rows.append(
                        Row(
                            number, number, number, dict(zip(fields, values, strict=True)), formulas
                        )
                    )
                    count += len(fields)
                    if count > MAX_CELLS or len(rows) > MAX_ROWS:
                        invalid("Use at most 100,000 workbook cells and 10,000 rows per sheet.")
                table = Table(sheet.title, fields, rows)
                check_table(table)
                tables.append(table)
            return Parsed(tables)
        finally:
            workbook.close()
    except DomainError:
        raise
    except Exception:
        # Library exceptions can contain spreadsheet cells or hostile XML. Never reflect them.
        invalid("The XLSX is damaged or unsafe. Save a plain XLSX workbook and try again.")


def parse(data: bytes, kind: str, options: ParseOptions) -> Parsed:
    if not data or len(data) > MAX_UPLOAD:
        invalid("Upload a nonempty file no larger than 8 MiB.")
    if kind == "csv":
        return parse_csv(data, options)
    if kind == "xlsx":
        return parse_xlsx(data)
    if kind == "md":
        text = decode(data, "utf-8-sig")
        if not text.strip():
            invalid("Markdown must contain nonempty UTF-8 text.")
        return Parsed([], text)
    invalid("Choose a .csv, .xlsx or .md file.")


def candidate_role(name: str) -> str:
    words = re.sub(r"[_\W]+", " ", name.casefold())
    if re.search(
        r"\b(outcome|revenue|profit|clicks|views|sales|resolved|paid|after|actual|result)\b", words
    ):
        return "outcome"
    if re.search(r"\b(target|ideal|label|advice|recommendation)\b", words):
        return "target"
    return "input"


def likely_type(value: str) -> str:
    if value.casefold() in {"true", "false"}:
        return "boolean"
    if re.match(r"^\d{4}-\d{2}-\d{2}(T| |$)", value):
        try:
            datetime.fromisoformat(value)
            return "date"
        except ValueError:
            pass
    # Ambiguous day/month ordering stays a date candidate, never converted.
    if re.fullmatch(r"\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}", value):
        return "date"
    try:
        if math.isfinite(float(value)):
            return "number"
    except ValueError:
        pass
    return "text"


def preview(table: Table) -> SheetPreview:
    summaries = []
    for field in table.fields:
        values = [row.values[field] for row in table.rows]
        present = [v for v in values if v is not None and v.strip()]
        types = {likely_type(v) for v in present}
        summaries.append(
            FieldSummary(
                name=field,
                missing=len(values) - len(present),
                likely_type=next(iter(types)) if len(types) == 1 else "mixed" if types else "empty",
                date_candidate="date" in types,
                candidate_role=candidate_role(field),
                formulas=sum(field in row.formulas for row in table.rows),
            )
        )
    return SheetPreview(
        name=table.name,
        fields=summaries,
        row_count=len(table.rows),
        samples=[{"row": row.number, "values": row.values} for row in table.rows[:5]],
    )


def markdown_sections(text: str) -> list[tuple[int, str]]:
    sections = [(0, "")]
    offset = 0
    fence = None
    lines = io.StringIO(text, newline="")
    previous = ""
    previous_offset = 0
    for line in lines:
        marker = re.match(r"^ {0,3}(`{3,}|~{3,})", line)
        if marker:
            if fence is None:
                fence = marker[1]
            elif marker[1][0] == fence[0] and len(marker[1]) >= len(fence):
                fence = None
        if fence is None and marker is None:
            heading = re.match(r"^ {0,3}#{1,6}\s+(.+?)\s*#*\s*$", line)
            if heading:
                sections.append((offset, heading[1][:200]))
            elif previous.strip() and re.fullmatch(r" {0,3}(?:=+|-+)\s*", line):
                sections.append((previous_offset, previous.strip()[:200]))
            if len(sections) > MAX_HEADINGS + 1:
                invalid("Split Markdown documents into at most 2,000 headings per file.")
            previous = line if heading is None else ""
        else:
            previous = ""
        previous_offset = offset
        offset += len(line)
    # A heading at offset zero replaces the empty preamble.
    return list(dict(sections).items())


def chunk_markdown(source_id: UUID, source_hash: str, text: str) -> list[Chunk]:
    result = []
    sections = markdown_sections(text)
    newlines = newline_positions(text)
    for index, (start, heading) in enumerate(sections):
        end = sections[index + 1][0] if index + 1 < len(sections) else len(text)
        while start < end:
            stop = min(start + CHUNK_LENGTH, end)
            if stop < end:
                boundary = text.rfind("\n", start, stop)
                if boundary >= start + CHUNK_LENGTH // 2:
                    stop = boundary + 1
            excerpt = text[start:stop]
            if excerpt.strip():
                digest = hashlib.sha256(excerpt.encode()).hexdigest()
                result.append(
                    Chunk(
                        citation_id=uuid5(NAMESPACE_URL, f"{source_id}:{start}:{stop}:{digest}"),
                        source_id=source_id,
                        source_sha256=source_hash,
                        heading=heading,
                        start=start,
                        end=stop,
                        line_start=bisect.bisect_left(newlines, start) + 1,
                        line_end=bisect.bisect_left(newlines, stop - 1) + 1,
                        text=excerpt,
                        sha256=digest,
                    )
                )
            start = stop
    return result


def newline_positions(text: str):
    # A compact array avoids one Python integer per newline in large text files.
    return array("I", (match.end() - 1 for match in re.finditer(r"\r\n|\r|\n", text)))
