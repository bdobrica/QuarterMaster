import hashlib
import io
import zipfile
from datetime import datetime
from uuid import uuid4

import pytest
from openpyxl import Workbook
from openpyxl.xml import DEFUSEDXML

from quartermaster.contracts.sources import ParseOptions
from quartermaster.errors import DomainError
from quartermaster.ingestion import (
    CHUNK_LENGTH,
    MAX_EXPANDED,
    chunk_markdown,
    parse,
    preview,
    safe_name,
)


def workbook_bytes():
    workbook = Workbook()
    first = workbook.active
    first.title = "Requests"
    first.append(["tool", "inspection_date", "paid_after", "note"])
    first.append(["drill", datetime(2026, 1, 2), 12, "<script>untrusted</script>"])
    first.append(["saw", None, "=1+1", '=HYPERLINK("https://invalid.test", "x")'])
    empty = workbook.create_sheet("Empty")
    assert empty.title == "Empty"
    second = workbook.create_sheet("Details")
    second.append(["tool", "status"])
    second.append(["drill", "ready"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def rewrite_zip(data, overrides):
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as original, zipfile.ZipFile(buffer, "w") as target:
        for info in original.infolist():
            target.writestr(info.filename, overrides.get(info.filename, original.read(info)))
        for name, content in overrides.items():
            if name not in original.namelist():
                target.writestr(name, content)
    return buffer.getvalue()


def test_csv_preserves_text_and_physical_multiline_positions():
    data = b' name ;date;paid_after;note\r\nA;01/02/2026;12;"first\r\nsecond"\r\nB;;0;=1+1\r\n'
    parsed = parse(data, "csv", ParseOptions(delimiter=";"))
    table = parsed.tables[0]
    assert table.fields == [" name ", "date", "paid_after", "note"]
    assert table.rows[0].values["date"] == "01/02/2026"
    assert table.rows[0].values["note"] == "first\r\nsecond"
    assert (table.rows[0].line_start, table.rows[0].line_end) == (2, 3)
    assert (table.rows[1].number, table.rows[1].line_start) == (3, 4)
    assert table.rows[1].values["note"] == "=1+1"
    fields = {f.name: f for f in preview(table).fields}
    assert fields["date"].missing == 1
    assert fields["date"].date_candidate
    assert fields["paid_after"].candidate_role == "outcome"
    assert fields["paid_after"].likely_type == "number"
    assert fields["note"].formulas == 0


def test_explicit_encoding_and_delimiter():
    data = "tool\tnote\ndrill\tcafé\n".encode("cp1252")
    with pytest.raises(DomainError, match="encoding"):
        parse(data, "csv", ParseOptions())
    parsed = parse(data, "csv", ParseOptions(encoding="cp1252", delimiter="\t"))
    assert parsed.tables[0].rows[0].values == {"tool": "drill", "note": "café"}


@pytest.mark.parametrize(
    "data",
    [b"", b"\x00\x01", b"a,a\nx,y", b"\na\n", b"a,b\nx", b'a\n"unterminated', b"a\n" + b"x" * 8001],
)
def test_csv_rejects_malformed_or_unsupported_content(data):
    with pytest.raises(DomainError, match="SOURCE_INVALID"):
        parse(data, "csv", ParseOptions())


def test_csv_row_and_cell_bounds():
    with pytest.raises(DomainError):
        parse(b"field\n" + b"row\n" * 10001, "csv", ParseOptions())
    with pytest.raises(DomainError):
        parse(
            ("a,b,c,d,e,f,g,h,i,j,k\n" + "1,2,3,4,5,6,7,8,9,10,11\n" * 9100).encode(),
            "csv",
            ParseOptions(),
        )
    with pytest.raises(DomainError):
        parse((",".join(f"f{i}" for i in range(101)) + "\n").encode(), "csv", ParseOptions())


def test_xlsx_multiple_sheets_dates_and_unevaluated_formulas():
    assert DEFUSEDXML is True
    data = workbook_bytes()
    parsed = parse(data, "xlsx", ParseOptions())
    assert [t.name for t in parsed.tables] == ["Requests", "Empty", "Details"]
    assert parsed.tables[0].rows[0].values["inspection_date"] == "2026-01-02T00:00:00"
    assert parsed.tables[0].rows[1].values["paid_after"] == "=1+1"
    fields = {f.name: f for f in preview(parsed.tables[0]).fields}
    assert fields["paid_after"].formulas == 1
    assert fields["inspection_date"].missing == 1
    assert fields["inspection_date"].date_candidate


@pytest.mark.parametrize(
    "overrides",
    [
        {"xl/vbaProject.bin": b"not executable"},
        {"../escape.xml": b"<a/>"},
        {
            "xl/worksheets/sheet1.xml": b'<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///etc/passwd">]><x>&secret;</x>'
        },
        {"xl/worksheets/sheet1.xml": b"<broken"},
        {"xl/worksheets/sheet1.xml": b'<worksheet><c r="XFD1048576"/></worksheet>'},
    ],
)
def test_xlsx_rejects_hostile_archives_and_xml(overrides):
    with pytest.raises(DomainError, match="SOURCE_INVALID"):
        parse(rewrite_zip(workbook_bytes(), overrides), "xlsx", ParseOptions())


def test_zip_expansion_limit_before_parsing():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("bomb.xml", b" " * (MAX_EXPANDED + 1))
    with pytest.raises(DomainError, match="expanded"):
        parse(buffer.getvalue(), "xlsx", ParseOptions())


def test_broken_xlsx_is_sanitized():
    with pytest.raises(DomainError) as caught:
        parse(b"private-cell-value", "xlsx", ParseOptions())
    assert "private-cell-value" not in str(caught.value)


@pytest.mark.parametrize(
    "name,expected",
    [
        ("../../source.csv", "source.csv"),
        ("C:\\private\\source.XLSX", "source.XLSX"),
        ("evil\r\nname.md", "evilname.md"),
        ("\u202eevil.md", "evil.md"),
        ("<script>.md", "<script>.md"),
    ],
)
def test_names_are_display_only_and_strip_paths_and_controls(name, expected):
    assert safe_name(name)[0] == expected


@pytest.mark.parametrize("name", ["run.py", "book.xlsm", "file.csv.exe", "README", "doc.txt"])
def test_unsupported_extensions(name):
    with pytest.raises(DomainError):
        safe_name(name)


def test_markdown_chunks_exact_positions_stable_ids_and_heading_boundaries():
    source = uuid4()
    text = (
        "Preamble\r\n# Safety\r\nCurrent inspection required.\r\n"
        + ("long unicode café\n" * 180)
        + "\nSecond heading\n---\nAsk first.\n```md\n# Not a heading\n```\n"
    )
    sha = hashlib.sha256(text.encode()).hexdigest()
    chunks = chunk_markdown(source, sha, text)
    assert chunks == chunk_markdown(source, sha, text)
    assert len(chunks) > 3
    assert "".join(c.text for c in chunks) == text
    assert {c.heading for c in chunks} == {"", "Safety", "Second heading"}
    for chunk in chunks:
        assert len(chunk.text) <= CHUNK_LENGTH
        assert text[chunk.start : chunk.end] == chunk.text
        assert chunk.sha256 == hashlib.sha256(chunk.text.encode()).hexdigest()
        assert chunk.line_start == text.count("\n", 0, chunk.start) + 1
        assert chunk.line_end == text.count("\n", 0, chunk.end - 1) + 1


def test_markdown_long_single_line_and_invalid_binary():
    text = "# Policy\n" + "a" * 4001
    assert all(len(c.text) <= CHUNK_LENGTH for c in chunk_markdown(uuid4(), "a" * 64, text))
    for data in (b"\x00binary", b"\xff", b"   "):
        with pytest.raises(DomainError):
            parse(data, "md", ParseOptions())


def test_markdown_does_not_treat_code_or_atx_followed_by_rule_as_setext():
    from quartermaster.ingestion import markdown_sections

    text = "# Heading\n---\n```\ntext\n```\n---\n"
    assert [h for _, h in markdown_sections(text)] == ["Heading"]


def test_markdown_heading_limit_is_actionable():
    from quartermaster.ingestion import markdown_sections

    with pytest.raises(DomainError, match="2,000 headings"):
        markdown_sections("# heading\n" * 2001)


@pytest.mark.parametrize("newline", ["\n", "\r\n", "\r"])
def test_markdown_source_line_positions_for_supported_newlines(newline):
    text = newline.join(["# First", "Evidence.", "# Second", "More evidence.", ""])
    chunks = chunk_markdown(uuid4(), hashlib.sha256(text.encode()).hexdigest(), text)
    assert [(c.line_start, c.line_end) for c in chunks] == [(1, 2), (3, 4)]
    assert "".join(c.text for c in chunks) == text
